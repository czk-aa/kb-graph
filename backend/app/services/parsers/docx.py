"""DOCX 文本提取（python-docx）。"""
from pathlib import Path

from docx import Document as DocxDocument


def parse(path: Path) -> str:
    doc = DocxDocument(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            style = (para.style.name or "").lower()
            if style.startswith("heading"):
                try:
                    level = int(style.removeprefix("heading").strip() or "1")
                except ValueError:
                    level = 2
                parts.append(f"{'#' * min(level, 6)} {para.text.strip()}")
            else:
                parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n\n".join(parts)
